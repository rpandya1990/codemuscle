import csv
import json
import shutil
import uuid
from datetime import date, datetime
from enum import Enum
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session, selectinload

from codemuscle.application.data_lifecycle.schemas import (
    BackupCreateRequest,
    BackupResponse,
    DeleteDataRequest,
    DeleteDataResponse,
    ExportRequest,
    ExportResponse,
    RestoreRequest,
    RestoreResponse,
)
from codemuscle.config import Settings
from codemuscle.domain.exceptions import (
    BackupNotFoundError,
    DataLifecycleError,
    WorkspaceNotInitializedError,
)
from codemuscle.infrastructure.database.base import Base
from codemuscle.infrastructure.database.models import (
    BackupRecord,
    Problem,
    UserPreference,
)

MANIFEST_VERSION = 1
APPLICATION_VERSION = "0.1.0"
TABLE_ORDER = [
    "import_jobs",
    "topics",
    "patterns",
    "problems",
    "problem_topics",
    "problem_patterns",
    "attempts",
    "import_rows",
    "queue_sessions",
    "queue_items",
    "user_preferences",
    "backup_records",
]


class DataLifecycleService:
    def __init__(self, session: Session, settings: Settings) -> None:
        self.session = session
        self.settings = settings

    def export(self, request: ExportRequest) -> ExportResponse:
        workspace = self._workspace()
        exports = workspace / "exports"
        exports.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        suffix = request.format.lower()
        filename = f"codemuscle-export-{timestamp}-{uuid.uuid4().hex[:8]}.{suffix}"
        target = exports / filename
        problems = list(
            self.session.scalars(
                select(Problem).options(
                    selectinload(Problem.topics),
                    selectinload(Problem.patterns),
                    selectinload(Problem.attempts),
                )
            ).unique()
        )
        if not request.include_archived:
            problems = [problem for problem in problems if problem.archived_at is None]
        if request.format == "CSV":
            self._write_csv(target, problems)
        elif request.format == "XLSX":
            self._write_xlsx(target, problems)
        else:
            target.write_text(
                json.dumps(self._export_payload(problems), indent=2, default=self._json_default),
                encoding="utf-8",
            )
        return ExportResponse(
            filename=filename,
            format=request.format,
            status="COMPLETED",
            problem_count=len(problems),
            attempt_count=sum(len(problem.attempts) for problem in problems),
        )

    def create_backup(self, request: BackupCreateRequest) -> BackupResponse:
        workspace = self._workspace()
        backups = workspace / "backups"
        backups.mkdir(parents=True, exist_ok=True)
        record = BackupRecord(
            filename="pending",
            manifest_version=MANIFEST_VERSION,
            application_version=APPLICATION_VERSION,
            status="CREATING",
        )
        self.session.add(record)
        self.session.flush()
        record.filename = f"codemuscle-backup-{datetime.now():%Y%m%d-%H%M%S}-{record.id}.zip"
        record.status = "COMPLETED"
        self.session.flush()
        manifest = {
            "manifest_version": MANIFEST_VERSION,
            "application_version": APPLICATION_VERSION,
            "created_at": datetime.now().isoformat(),
            "database_format": "codemuscle-json-v1",
            "included_directories": [
                name
                for name, included in {
                    "imports": request.include_imports,
                    "exports": request.include_exports,
                }.items()
                if included
            ],
        }
        target = backups / record.filename
        with ZipFile(target, "w", ZIP_DEFLATED) as archive:
            archive.writestr("manifest.json", json.dumps(manifest, indent=2))
            archive.writestr("database.json", json.dumps(self._snapshot(), indent=2))
            for directory in manifest["included_directories"]:
                source = workspace / directory
                if source.exists():
                    for file in source.rglob("*"):
                        if file.is_file():
                            archive.write(file, Path(directory) / file.relative_to(source))
        self.session.commit()
        self.session.refresh(record)
        return BackupResponse.model_validate(record)

    def list_backups(self) -> list[BackupResponse]:
        records = self.session.scalars(
            select(BackupRecord).order_by(BackupRecord.created_at.desc())
        ).all()
        return [BackupResponse.model_validate(record) for record in records]

    def restore(self, backup_id: uuid.UUID, request: RestoreRequest) -> RestoreResponse:
        if request.confirmation != "RESTORE":
            raise DataLifecycleError("Type RESTORE to confirm replacing current application data.")
        record = self.session.get(BackupRecord, backup_id)
        if record is None:
            raise BackupNotFoundError(backup_id)
        workspace = self._workspace()
        source = workspace / "backups" / record.filename
        if not source.is_file():
            raise DataLifecycleError("The backup archive is missing from the private workspace.")
        with ZipFile(source) as archive:
            names = set(archive.namelist())
            if not {"manifest.json", "database.json"}.issubset(names):
                raise DataLifecycleError(
                    "The backup archive is missing its manifest or database export."
                )
            manifest = json.loads(archive.read("manifest.json"))
            if manifest.get("manifest_version") != MANIFEST_VERSION:
                raise DataLifecycleError("The backup manifest version is not supported.")
            snapshot = json.loads(archive.read("database.json"))
            restored_rows = self._restore_snapshot(snapshot)
            for name in names:
                parts = Path(name).parts
                if len(parts) < 2 or parts[0] not in {"imports", "exports"} or name.endswith("/"):
                    continue
                destination = workspace.joinpath(*parts)
                destination.parent.mkdir(parents=True, exist_ok=True)
                destination.write_bytes(archive.read(name))
        self.session.commit()
        return RestoreResponse(
            backup_id=backup_id,
            status="COMPLETED",
            restored_tables=len(TABLE_ORDER),
            restored_rows=restored_rows,
        )

    def delete_all(self, request: DeleteDataRequest) -> DeleteDataResponse:
        if request.confirmation != "DELETE ALL DATA":
            raise DataLifecycleError("Type DELETE ALL DATA to confirm permanent deletion.")
        workspace = self._workspace()
        cleared: list[str] = []
        for name, enabled in {
            "imports": request.delete_import_files,
            "exports": request.delete_export_files,
            "backups": request.delete_backup_files,
        }.items():
            if enabled:
                self._clear_directory(workspace / name)
                cleared.append(name)
        deleted_rows = self._delete_tables(include_backups=request.delete_backup_files)
        self.session.commit()
        return DeleteDataResponse(
            status="COMPLETED",
            deleted_rows=deleted_rows,
            cleared_directories=cleared,
            backups_preserved=not request.delete_backup_files,
        )

    def _workspace(self) -> Path:
        preference = self.session.scalar(select(UserPreference).limit(1))
        value = preference.workspace_path if preference else self.settings.workspace_path
        if value is None:
            raise WorkspaceNotInitializedError()
        return Path(value)

    def _snapshot(self) -> dict[str, list[dict[str, object]]]:
        snapshot: dict[str, list[dict[str, object]]] = {}
        for name in TABLE_ORDER:
            table = Base.metadata.tables[name]
            rows = self.session.execute(select(table)).mappings().all()
            snapshot[name] = [
                {key: self._encode(value) for key, value in row.items()} for row in rows
            ]
        return snapshot

    def _restore_snapshot(self, snapshot: dict[str, Any]) -> int:
        if not isinstance(snapshot, dict) or any(name not in snapshot for name in TABLE_ORDER):
            raise DataLifecycleError("The database export is incomplete.")
        self._delete_tables(include_backups=True)
        restored = 0
        for name in TABLE_ORDER:
            table = Base.metadata.tables[name]
            rows = snapshot[name]
            if not isinstance(rows, list):
                raise DataLifecycleError(f"Invalid rows for table {name}.")
            decoded = [
                {
                    column.name: self._decode(row.get(column.name), column.type)
                    for column in table.columns
                }
                for row in rows
            ]
            if decoded:
                self.session.execute(table.insert(), decoded)
                restored += len(decoded)
        return restored

    def _delete_tables(self, *, include_backups: bool) -> int:
        deleted = 0
        for name in reversed(TABLE_ORDER):
            if name == "backup_records" and not include_backups:
                continue
            table = Base.metadata.tables[name]
            deleted += self.session.scalar(select(func.count()).select_from(table)) or 0
            self.session.execute(delete(table))
        return deleted

    @staticmethod
    def _encode(value: object) -> object:
        if isinstance(value, Enum):
            return {"$type": "enum", "value": value.value}
        if isinstance(value, uuid.UUID):
            return {"$type": "uuid", "value": str(value)}
        if isinstance(value, datetime):
            return {"$type": "datetime", "value": value.isoformat()}
        if isinstance(value, date):
            return {"$type": "date", "value": value.isoformat()}
        return value

    @staticmethod
    def _decode(value: object, column_type: Any) -> object:
        if not isinstance(value, dict) or "$type" not in value:
            return value
        kind, raw = value["$type"], value["value"]
        if kind == "uuid":
            return uuid.UUID(raw)
        if kind == "datetime":
            return datetime.fromisoformat(raw)
        if kind == "date":
            return date.fromisoformat(raw)
        if kind == "enum" and getattr(column_type, "enum_class", None):
            return column_type.enum_class(raw)
        return raw

    @staticmethod
    def _json_default(value: object) -> str:
        if isinstance(value, (date, datetime, uuid.UUID, Enum)):
            return str(value.value if isinstance(value, Enum) else value)
        raise TypeError(f"Cannot serialize {type(value).__name__}")

    @staticmethod
    def _export_payload(problems: list[Problem]) -> dict[str, object]:
        return {
            "export_version": 1,
            "exported_at": datetime.now().isoformat(),
            "problems": [
                {
                    "id": problem.id,
                    "title": problem.title,
                    "url": problem.url,
                    "difficulty": problem.difficulty,
                    "notes": problem.notes,
                    "priority": problem.priority,
                    "mastery": problem.current_mastery_state,
                    "next_revision_date": problem.next_revision_date,
                    "topics": [topic.name for topic in problem.topics],
                    "patterns": [pattern.name for pattern in problem.patterns],
                    "attempts": [
                        {
                            "attempted_at": attempt.attempted_at,
                            "outcome": attempt.outcome,
                            "hint_usage": attempt.hint_usage,
                            "time_spent_minutes": attempt.time_spent_minutes,
                            "notes": attempt.notes,
                            "schedule_explanation": attempt.schedule_explanation,
                        }
                        for attempt in problem.attempts
                    ],
                }
                for problem in problems
            ],
        }

    @staticmethod
    def _write_csv(target: Path, problems: list[Problem]) -> None:
        with target.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=[
                    "title",
                    "url",
                    "difficulty",
                    "notes",
                    "priority",
                    "mastery",
                    "next_revision_date",
                    "topics",
                    "patterns",
                    "attempt_count",
                ],
            )
            writer.writeheader()
            for problem in problems:
                writer.writerow(
                    {
                        "title": problem.title,
                        "url": problem.url or "",
                        "difficulty": problem.difficulty.value,
                        "notes": problem.notes or "",
                        "priority": problem.priority,
                        "mastery": problem.current_mastery_state.value,
                        "next_revision_date": problem.next_revision_date or "",
                        "topics": ", ".join(topic.name for topic in problem.topics),
                        "patterns": ", ".join(pattern.name for pattern in problem.patterns),
                        "attempt_count": len(problem.attempts),
                    }
                )

    @staticmethod
    def _write_xlsx(target: Path, problems: list[Problem]) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Problems"
        sheet.append(
            [
                "Title",
                "URL",
                "Difficulty",
                "Notes",
                "Priority",
                "Mastery",
                "Next revision",
                "Topics",
                "Patterns",
            ]
        )
        attempts = workbook.create_sheet("Attempts")
        attempts.append(
            [
                "Problem",
                "Attempted at",
                "Outcome",
                "Hint usage",
                "Minutes",
                "Notes",
                "Explanation",
            ]
        )
        for problem in problems:
            sheet.append(
                [
                    problem.title,
                    problem.url,
                    problem.difficulty.value,
                    problem.notes,
                    problem.priority,
                    problem.current_mastery_state.value,
                    problem.next_revision_date,
                    ", ".join(topic.name for topic in problem.topics),
                    ", ".join(pattern.name for pattern in problem.patterns),
                ]
            )
            for attempt in problem.attempts:
                attempts.append(
                    [
                        problem.title,
                        attempt.attempted_at.replace(tzinfo=None),
                        attempt.outcome.value,
                        attempt.hint_usage.value,
                        attempt.time_spent_minutes,
                        attempt.notes,
                        attempt.schedule_explanation,
                    ]
                )
        workbook.save(target)

    @staticmethod
    def _clear_directory(directory: Path) -> None:
        directory.mkdir(parents=True, exist_ok=True)
        for child in directory.iterdir():
            if child.is_dir():
                shutil.rmtree(child)
            else:
                child.unlink()
